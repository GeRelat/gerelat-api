"""
renderer.py — ProseMirror JSON → python-docx Document → BytesIO
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt, RGBColor

# ── Inline helpers ────────────────────────────────────────────────


def _apply_marks(run: Any, marks: list[dict]) -> None:
    for mark in marks:
        t = mark.get("type")
        if t in ("bold", "strong"):
            run.bold = True
        elif t in ("italic", "em"):
            run.italic = True
        elif t == "underline":
            run.underline = True
        elif t == "code":
            run.font.name = "Courier New"
            run.font.size = Pt(10)


def _fill_inline(para: Any, nodes: list[dict]) -> None:
    for node in nodes:
        if node.get("type") == "text":
            run = para.add_run(node.get("text", ""))
            _apply_marks(run, node.get("marks", []))
        elif node.get("type") == "hard_break":
            para.add_run().add_break()


# ── Block renderers ───────────────────────────────────────────────


def _render_paragraph(doc: Any, node: dict) -> None:
    para = doc.add_paragraph()
    _fill_inline(para, node.get("content", []))


_HEADING_SIZES = {1: Pt(18), 2: Pt(14), 3: Pt(12)}


def _render_heading(doc: Any, node: dict) -> None:
    level = node.get("attrs", {}).get("level", 1)
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    para.paragraph_format.space_after = Pt(6)
    _fill_inline(para, node.get("content", []))
    # Aplica tamanho e negrito a todos os runs do heading
    for r in para.runs:
        r.bold = True
        r.font.size = _HEADING_SIZES.get(level, Pt(11))
        r.font.color.rgb = RGBColor(0, 0, 0)  # preto — sem azul do estilo padrão


def _render_blockquote(doc: Any, node: dict) -> None:
    for child in node.get("content", []):
        if child.get("type") == "paragraph":
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Mm(10)
            _fill_inline(para, child.get("content", []))


def _render_list(doc: Any, node: dict) -> None:
    is_bullet = node.get("type") == "bullet_list"
    style_name = "List Bullet" if is_bullet else "List Number"
    for item in node.get("content", []):
        for block in item.get("content", []):
            para = doc.add_paragraph(style=style_name)
            _fill_inline(para, block.get("content", []))


def _render_data_table(doc: Any, node: dict) -> None:
    import json as _json

    attrs = node.get("attrs", {})
    # columns/rows podem estar em attrs (JSON strings) ou diretos no nó (arrays)
    raw_cols = attrs.get("columns", node.get("columns", []))
    raw_rows = attrs.get("rows", node.get("rows", []))
    columns: list[dict] = (
        _json.loads(raw_cols) if isinstance(raw_cols, str) else raw_cols
    )
    rows: list[dict] = _json.loads(raw_rows) if isinstance(raw_rows, str) else raw_rows
    caption: str = attrs.get("caption", "")

    if not columns:
        return

    if caption:
        doc.add_paragraph(caption)

    table = doc.add_table(rows=1 + len(rows), cols=len(columns))
    table.style = "Table Grid"

    # Header row — bold labels
    hdr = table.rows[0]
    for j, col in enumerate(columns):
        cell = hdr.cells[j]
        cell.text = col.get("label", col.get("key", ""))
        for run in cell.paragraphs[0].runs:
            run.bold = True

    # Data rows
    for i, row_data in enumerate(rows):
        row = table.rows[i + 1]
        for j, col in enumerate(columns):
            row.cells[j].text = str(row_data.get(col.get("key", ""), ""))


def _render_custom_text(doc: Any, node: dict) -> None:
    """Renderiza bloco editável — itera seus blocos filhos normalmente."""
    for child in node.get("content", []):
        _render_block(doc, child)


# Nós de variável (bloco) que chegam aqui só quando não foram resolvidos —
# variável sem valor. Renderiza o placeholder [ref] em vez de sumir com ele.
_UNRESOLVED_BLOCK_VAR_TYPES = {
    "var_richtext",
    "var_image",
    "image_ref",
    "var_table",
    "table_ref",
}


def _render_unresolved_var(doc: Any, node: dict) -> None:
    ref = node.get("attrs", {}).get("ref", "")
    doc.add_paragraph(f"[{ref}]" if ref else "[variável não resolvida]")


def _render_block(doc: Any, node: dict) -> None:
    node_type = node.get("type")
    if node_type == "paragraph":
        _render_paragraph(doc, node)
    elif node_type == "heading":
        _render_heading(doc, node)
    elif node_type == "blockquote":
        _render_blockquote(doc, node)
    elif node_type in ("bullet_list", "ordered_list"):
        _render_list(doc, node)
    elif node_type == "data_table":
        _render_data_table(doc, node)
    elif node_type == "image_block":
        _render_image_block(doc, node)
    elif node_type == "custom_text":
        _render_custom_text(doc, node)
    elif node_type in _UNRESOLVED_BLOCK_VAR_TYPES:
        _render_unresolved_var(doc, node)
    # Outros tipos são silenciosamente ignorados


def _abrir_asset_stream(asset_id: Any) -> BytesIO | None:
    """Abre o arquivo de um Asset como BytesIO (seekable, reutilizável), ou None
    se o id for vazio ou o arquivo não puder ser lido."""
    if not asset_id:
        return None
    try:
        from django.apps import apps
        from django.core.files.storage import default_storage

        Asset = apps.get_model("gerelat", "Asset")
        asset = Asset.objects.get(pk=asset_id)
        with default_storage.open(str(asset.arquivo), "rb") as f:
            return BytesIO(f.read())
    except Exception:
        return None


def _render_image_block(doc: Any, node: dict) -> None:
    attrs = node.get("attrs", {})
    asset_id = attrs.get("asset_id")
    width_mm = attrs.get("width_mm")
    caption = attrs.get("caption", "")

    if not asset_id:
        doc.add_paragraph("[imagem não encontrada]")
        return

    stream = _abrir_asset_stream(asset_id)
    if stream is None:
        doc.add_paragraph(f"[erro ao carregar imagem #{asset_id}]")
        return

    width = Mm(float(width_mm)) if width_mm else None
    doc.add_picture(stream, width=width)
    if caption:
        doc.add_paragraph(caption)


# ── Cabeçalho com logomarcas ─────────────────────────────────────

# Altura alvo das logomarcas no cabeçalho — mesma altura nas duas, largura
# proporcional, para alinharem independentemente da proporção de cada arquivo.
_LOGO_ALTURA_MM = 16


def logos_cabecalho_de_metadados(metadados: dict) -> dict:
    """Extrai os ids das logomarcas de cliente/parceira congelados em
    `Relatorio.metadados` (ver `build_variables_context`)."""
    metadados = metadados or {}
    return {
        "logo_cliente_id": metadados.get("sistema.cabecalho.logo_cliente_asset_id"),
        "logo_parceira_id": metadados.get("sistema.cabecalho.logo_parceira_asset_id"),
    }


def _render_cabecalho_logos(
    doc: Any,
    *,
    logo_cliente_id: Any = None,
    logo_parceira_id: Any = None,
) -> None:
    """Cabeçalho de página: logomarca do cliente à esquerda, da parceira à
    direita. Omitido inteiro se nenhuma das duas existir."""
    s_cliente = _abrir_asset_stream(logo_cliente_id)
    s_parceira = _abrir_asset_stream(logo_parceira_id)
    if s_cliente is None and s_parceira is None:
        return

    section = doc.sections[0]
    header = section.header
    header.is_linked_to_previous = False
    header.paragraphs[0].text = ""

    largura = section.page_width - section.left_margin - section.right_margin
    tabela = header.add_table(rows=1, cols=2, width=largura)
    celula_esq, celula_dir = tabela.rows[0].cells

    if s_cliente is not None:
        p = celula_esq.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.add_run().add_picture(s_cliente, height=Mm(_LOGO_ALTURA_MM))
    if s_parceira is not None:
        p = celula_dir.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.add_run().add_picture(s_parceira, height=Mm(_LOGO_ALTURA_MM))


# ── Main entry point ──────────────────────────────────────────────


def prosemirror_to_docx(
    pm_json: dict,
    *,
    logo_cliente_id: Any = None,
    logo_parceira_id: Any = None,
) -> BytesIO:
    """
    Convert a resolved ProseMirror JSON document (after placeholder replacement)
    to a python-docx Document and return it as a BytesIO buffer.

    `logo_cliente_id` / `logo_parceira_id`: ids de Asset opcionais para o
    cabeçalho de página (cliente à esquerda, parceira à direita).
    """
    doc = Document()

    _render_cabecalho_logos(
        doc, logo_cliente_id=logo_cliente_id, logo_parceira_id=logo_parceira_id
    )

    for node in pm_json.get("content", []):
        _render_block(doc, node)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
