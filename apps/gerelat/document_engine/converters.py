"""
converters.py
  - convert_docx_to_prosemirror: reads a .docx file, returns ProseMirror JSON
  - replace_placeholders: walks a ProseMirror template JSON and substitutes
    variable/image_ref/table_ref/custom_text nodes with actual campo values
"""

from __future__ import annotations

import copy
import re
from typing import Any

from docx import Document
from docx.oxml.ns import qn

# ── DOCX → ProseMirror ────────────────────────────────────────────


def _heading_level(style_val: str) -> int | None:
    """Return 1-3 if style_val is a heading style id, else None."""
    m = re.match(r"heading(\d)", style_val, re.IGNORECASE)
    if m:
        lvl = int(m.group(1))
        return lvl if 1 <= lvl <= 3 else None
    return None


def _parse_run(r_elem: Any) -> list[dict]:
    """Convert a w:r XML element to a list of ProseMirror inline nodes."""
    texts = [t.text or "" for t in r_elem.findall(f".//{qn('w:t')}")]
    text = "".join(texts)

    if not text:
        if r_elem.find(f".//{qn('w:br')}") is not None:
            return [{"type": "hard_break"}]
        return []

    marks: list[dict] = []
    rpr = r_elem.find(qn("w:rPr"))
    if rpr is not None:
        b = rpr.find(qn("w:b"))
        if b is not None and b.get(qn("w:val"), "true") not in ("false", "0"):
            marks.append({"type": "strong"})
        i_el = rpr.find(qn("w:i"))
        if i_el is not None and i_el.get(qn("w:val"), "true") not in ("false", "0"):
            marks.append({"type": "em"})
        u_el = rpr.find(qn("w:u"))
        if u_el is not None and u_el.get(qn("w:val"), "none") not in (
            "none",
            "false",
            "0",
        ):
            marks.append({"type": "underline"})

    node: dict = {"type": "text", "text": text}
    if marks:
        node["marks"] = marks
    return [node]


def _parse_inline(p_elem: Any) -> list[dict]:
    """Extract all inline nodes from a w:p or w:hyperlink element."""
    nodes: list[dict] = []
    for child in p_elem:
        tag = child.tag
        if tag == qn("w:r"):
            nodes.extend(_parse_run(child))
        elif tag == qn("w:hyperlink"):
            for r in child.findall(f".//{qn('w:r')}"):
                nodes.extend(_parse_run(r))
    return nodes


def _parse_paragraph(p_elem: Any) -> dict | None:
    ppr = p_elem.find(qn("w:pPr"))
    style_val = ""
    if ppr is not None:
        pstyle = ppr.find(qn("w:pStyle"))
        if pstyle is not None:
            style_val = pstyle.get(qn("w:val"), "")

    inline = _parse_inline(p_elem)
    level = _heading_level(style_val)

    if level:
        return {"type": "heading", "attrs": {"level": level}, "content": inline}
    return {"type": "paragraph", "content": inline}


def _parse_table(tbl_elem: Any) -> dict:
    columns: list[dict] | None = None
    rows_data: list[dict] = []

    for i, tr in enumerate(tbl_elem.findall(qn("w:tr"))):
        cells: list[str] = []
        for tc in tr.findall(qn("w:tc")):
            parts: list[str] = []
            for p in tc.findall(f".//{qn('w:p')}"):
                text = "".join(
                    t.text or ""
                    for r in p.findall(f".//{qn('w:r')}")
                    for t in r.findall(f".//{qn('w:t')}")
                )
                if text:
                    parts.append(text)
            cells.append(" ".join(parts).strip())

        if i == 0:
            columns = [
                {
                    "key": re.sub(r"\W+", "_", c).lower() or f"col{j}",
                    "label": c or f"Col{j}",
                }
                for j, c in enumerate(cells)
            ]
        else:
            if columns:
                rows_data.append({col["key"]: val for col, val in zip(columns, cells)})

    return {
        "type": "data_table",
        "attrs": {"caption": ""},
        "columns": columns or [],
        "rows": rows_data,
    }


def convert_docx_to_prosemirror(docx_file: Any) -> dict:
    """
    Read a .docx file and return a ProseMirror JSON document.
    Preserves bold/italic/underline, headings, and tables.
    """
    doc = Document(docx_file)
    content: list[dict] = []

    for child in doc.element.body:
        tag = child.tag
        if tag == qn("w:p"):
            node = _parse_paragraph(child)
            if node is not None:
                content.append(node)
        elif tag == qn("w:tbl"):
            content.append(_parse_table(child))

    if not content:
        content = [{"type": "paragraph", "content": []}]

    return {"type": "doc", "content": content}


# ── Placeholder replacement ───────────────────────────────────────

_MARK_ALIASES = {"bold": "strong", "italic": "em"}


def _normalize_marks(node: dict) -> dict:
    """Recursively rename legacy mark types (bold→strong, italic→em)."""
    node = dict(node)
    if "marks" in node:
        node["marks"] = [
            {**m, "type": _MARK_ALIASES.get(m.get("type", ""), m.get("type", ""))}
            for m in node["marks"]
        ]
    if "content" in node:
        node["content"] = [_normalize_marks(child) for child in node["content"]]
    return node


# Profundidade máxima de aninhamento de variáveis richtext (uma variável richtext
# resolvida cujo valor referencia outra variável richtext, e assim por diante).
# `validate_richtext` já proíbe richtext-dentro-de-richtext, mas documentos
# legados podem ter o nó — a guarda evita `RecursionError` em produção.
MAX_PROFUNDIDADE_RICHTEXT = 5


def _paragrafo_texto(texto: str) -> dict:
    return {"type": "paragraph", "content": [{"type": "text", "text": texto}]}


def _inline_nodes_from_valor(
    ref: str,
    valor: Any,
    variables: dict,
    pilha: frozenset,
    profundidade: int,
) -> list[dict]:
    """Return inline ProseMirror nodes for a text/richtext campo valor.

    Um valor richtext pode conter outras variáveis (var_text/var_image/...) —
    depois de achatar seus blocos em nós inline, resolve-as recursivamente."""
    if not valor:
        return [{"type": "text", "text": f"[{ref}]"}]

    if isinstance(valor, dict) and valor.get("type") == "doc":
        inline: list[dict] = []
        for block in valor.get("content", []):
            inline.extend(block.get("content") or [])
        inline = _traverse_inline(
            inline,
            variables,
            pilha=pilha | {ref},
            profundidade=profundidade + 1,
        )
        return [_normalize_marks(node) for node in inline] or [
            {"type": "text", "text": ""}
        ]

    return [{"type": "text", "text": str(valor)}]


def _block_nodes_for_image_ref(ref: str, node_attrs: dict, valor: Any) -> list[dict]:
    if not isinstance(valor, dict) or "asset_id" not in valor:
        return [
            {
                "type": "paragraph",
                "content": [{"type": "text", "text": f"[imagem {ref} não encontrada]"}],
            }
        ]
    return [
        {
            "type": "image_block",
            "attrs": {
                "asset_id": valor.get("asset_id"),
                "width_mm": valor.get("width_mm"),
                "height_mm": valor.get("height_mm"),
                "caption": valor.get("caption") or node_attrs.get("caption", ""),
            },
        }
    ]


def _block_nodes_for_table_ref(ref: str, node_attrs: dict, valor: Any) -> list[dict]:
    if not isinstance(valor, dict) or "columns" not in valor:
        return [
            {
                "type": "paragraph",
                "content": [{"type": "text", "text": f"[tabela {ref} não encontrada]"}],
            }
        ]
    import json as _json

    return [
        {
            "type": "data_table",
            "attrs": {
                "caption": node_attrs.get("caption", ""),
                "columns": _json.dumps(valor.get("columns", []), ensure_ascii=False),
                "rows": _json.dumps(valor.get("rows", []), ensure_ascii=False),
            },
        }
    ]


def _block_nodes_for_custom_text(
    label: str, existing_content: list[dict], valor: Any
) -> list[dict]:
    """Use richtext campo valor if available, else keep existing custom_text content."""
    if isinstance(valor, dict) and valor.get("type") == "doc":
        return [_normalize_marks(b) for b in valor.get("content", [])]
    return existing_content


_INLINE_VAR_TYPES = {"var_text", "variable"}
_IMAGE_VAR_TYPES = {"var_image", "image_ref"}
_TABLE_VAR_TYPES = {"var_table", "table_ref"}
_RICHTEXT_VAR_TYPES = {"var_richtext"}


def _traverse_inline(
    nodes: list[dict],
    variables: dict,
    *,
    pilha: frozenset = frozenset(),
    profundidade: int = 0,
) -> list[dict]:
    result: list[dict] = []
    for node in nodes:
        if node.get("type") in _INLINE_VAR_TYPES:
            ref = node.get("attrs", {}).get("ref", "")
            if ref in pilha or profundidade >= MAX_PROFUNDIDADE_RICHTEXT:
                result.append(
                    {"type": "text", "text": f'[ciclo de variáveis em "{ref}"]'}
                )
                continue
            valor = variables.get(ref)
            result.extend(
                _inline_nodes_from_valor(ref, valor, variables, pilha, profundidade)
            )
        else:
            result.append(node)
    return result


def _traverse_blocks(
    nodes: list[dict],
    variables: dict,
    *,
    pilha: frozenset = frozenset(),
    profundidade: int = 0,
) -> list[dict]:
    result: list[dict] = []
    for node in nodes:
        node_type = node.get("type")
        attrs = node.get("attrs", {})

        if node_type in _IMAGE_VAR_TYPES:
            ref = attrs.get("ref", "")
            result.extend(_block_nodes_for_image_ref(ref, attrs, variables.get(ref)))

        elif node_type in _TABLE_VAR_TYPES:
            ref = attrs.get("ref", "")
            result.extend(_block_nodes_for_table_ref(ref, attrs, variables.get(ref)))

        elif node_type in _RICHTEXT_VAR_TYPES:
            # var_richtext: substitui pelo conteúdo richtext da variável, com as
            # variáveis aninhadas nesse conteúdo também resolvidas.
            ref = attrs.get("ref", "")
            valor = variables.get(ref)
            if ref in pilha or profundidade >= MAX_PROFUNDIDADE_RICHTEXT:
                result.append(
                    _paragrafo_texto(f'[ciclo de variáveis em "{ref}"]')
                )
            elif isinstance(valor, dict) and valor.get("type") == "doc":
                filhos = _traverse_blocks(
                    valor.get("content", []),
                    variables,
                    pilha=pilha | {ref},
                    profundidade=profundidade + 1,
                )
                result.extend(_normalize_marks(b) for b in filhos)
            else:
                # Mantém como placeholder se a variável não tiver valor
                result.append(node)

        elif node_type == "custom_text":
            label = attrs.get("label", "")
            valor = variables.get(label)
            if (
                label not in pilha
                and profundidade < MAX_PROFUNDIDADE_RICHTEXT
                and isinstance(valor, dict)
                and valor.get("type") == "doc"
            ):
                # Backward compat: custom_text cujo label corresponde a uma variável richtext
                filhos = _traverse_blocks(
                    valor.get("content", []),
                    variables,
                    pilha=pilha | {label},
                    profundidade=profundidade + 1,
                )
                result.extend(_normalize_marks(b) for b in filhos)
            else:
                # Bloco editável livre — mantém como custom_text para edição no relatório
                new_node = copy.copy(node)
                new_node["content"] = _traverse_blocks(
                    node.get("content", []),
                    variables,
                    pilha=pilha,
                    profundidade=profundidade,
                )
                result.append(new_node)

        elif node_type == "paragraph":
            new_node = copy.copy(node)
            new_node["content"] = _traverse_inline(
                node.get("content", []),
                variables,
                pilha=pilha,
                profundidade=profundidade,
            )
            result.append(new_node)

        elif "content" in node:
            new_node = copy.copy(node)
            new_node["content"] = _traverse_blocks(
                node["content"], variables, pilha=pilha, profundidade=profundidade
            )
            result.append(new_node)

        else:
            result.append(node)

    return result


def extract_gap_refs(
    template_json: dict, docs_por_ref: dict[str, dict] | None = None
) -> set[str]:
    """Walk a ProseMirror template JSON and collect every ref/label referenced
    by a gap node (var_text/variable/var_richtext/var_image/image_ref/
    var_table/table_ref/custom_text).

    Quando `docs_por_ref` é fornecido (mapa referencia -> documento ProseMirror
    do valor de uma variável richtext), a varredura desce também para dentro
    desses valores, expondo as variáveis usadas transitivamente por eles."""
    refs: set[str] = set()
    gap_ref_types = (
        _INLINE_VAR_TYPES | _IMAGE_VAR_TYPES | _TABLE_VAR_TYPES | _RICHTEXT_VAR_TYPES
    )
    visitados: set[str] = set()

    def _walk(nodes: list[dict]) -> None:
        for node in nodes:
            node_type = node.get("type")
            attrs = node.get("attrs") or {}
            ref = None
            if node_type in gap_ref_types:
                ref = attrs.get("ref")
            elif node_type == "custom_text":
                ref = attrs.get("label")
            if ref:
                refs.add(ref)
                if docs_por_ref and ref not in visitados:
                    visitados.add(ref)
                    doc = docs_por_ref.get(ref)
                    if isinstance(doc, dict):
                        _walk(doc.get("content", []))
            if "content" in node:
                _walk(node["content"])

    _walk(template_json.get("content", []))
    return refs


def replace_placeholders(template_json: dict, variables: dict) -> dict:
    """
    Walk a ProseMirror template JSON and replace all placeholder nodes
    (variable, image_ref, table_ref, custom_text) with their resolved content.

    variables: {referencia: valor} from ProjetoCampo records
    """
    result = copy.deepcopy(template_json)
    result["content"] = _traverse_blocks(result.get("content", []), variables)
    return _normalize_marks(result)
