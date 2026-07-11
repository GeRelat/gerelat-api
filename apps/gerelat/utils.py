import os
import json
import boto3
import bleach
from botocore.exceptions import ClientError
from django.conf import settings
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def obter_imagem_real(relatorio, ref):
    try:
        partes = ref.split('.')
        entidade = partes[0]
        campo = partes[1] if len(partes) > 1 else None
        if entidade == 'evento' and relatorio.evento:
            imagem_field = getattr(relatorio.evento, campo, None)
            if imagem_field and hasattr(imagem_field, 'path'):
                return imagem_field.path
        elif entidade == 'relatorio':
            imagem_field = getattr(relatorio, campo, None)
            if imagem_field and hasattr(imagem_field, 'path'):
                return imagem_field.path

        return None
    except Exception:
        return None
    
def renderizar_imagem(node, document, relatorio):
    ref = node.get('attrs', {}).get('ref', '')
    caption = node.get('attrs', {}).get('caption', '')
    caminho_img = obter_imagem_real(relatorio, ref)
    
    if caminho_img and os.path.exists(caminho_img):
        try:
            p_img = document.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p_img.add_run()
            
            run.add_picture(caminho_img, width=Inches(5.5))
            if caption:
                p_cap = document.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run_cap = p_cap.add_run(caption)
                run_cap.italic = True
                run_cap.font.size = Pt(9)
                run_cap.font.color.rgb = RGBColor(80, 80, 80) 
        except Exception as e:
            document.add_paragraph(f"[ Erro ao processar imagem: {ref} ]")
    else:
        p_err = document.add_paragraph(f"[ Imagem não encontrada: {ref} ]")
        p_err.alignment = WD_ALIGN_PARAGRAPH.CENTER

def limpar_json_malicioso(dados):
    if isinstance(dados, dict):
        return {chave: limpar_json_malicioso(valor) for chave, valor in dados.items()}
    elif isinstance(dados, list):
        return [limpar_json_malicioso(item) for item in dados]
    elif isinstance(dados, str):
        return bleach.clean(dados, strip=True)
    else:
        return dados
    

def obter_valor_real(relatorio, referencia):
    try:
        partes = referencia.split('.')
        entidade = partes[0]
        campo = partes[1] if len(partes) > 1 else None

        if entidade == 'evento' and relatorio.evento:
            valor = getattr(relatorio.evento, campo, f"[{referencia}]")
            return str(valor)
        
        elif entidade == 'relatorio':
            valor = getattr(relatorio, campo, f"[{referencia}]")
            return str(valor)
            
        return f"[{referencia}]"
    except Exception:
        return f"[{referencia}]"

def obter_lista_real(relatorio, ref):
    try:
        partes = ref.split('.')
        entidade = partes[0] 

        if entidade == 'parceiros' and hasattr(relatorio, 'parceiros'):
            return relatorio.parceiros.all()
        
        elif hasattr(relatorio, 'evento') and relatorio.evento:
            if hasattr(relatorio.evento, entidade):
                relacao = getattr(relatorio.evento, entidade)
                if hasattr(relacao, 'all'):
                    return relacao.all()
        return []
    except Exception:
        return []
    
def renderizar_tabela(node, document, relatorio):
    ref = node.get('attrs', {}).get('ref', '')
    columns_str = node.get('attrs', {}).get('columns', '[]')
    
    try:
        import json
        columns = json.loads(columns_str)
    except:
        columns = ['Coluna 1']
        
    if not columns:
        columns = ['C1']
        
    lista_dados = obter_lista_real(relatorio, ref)
    
    table = document.add_table(rows=1, cols=len(columns))
    table.style = 'Table Grid' 
    
    hdr_cells = table.rows[0].cells
    for i, col in enumerate(columns):
        hdr_cells[i].text = str(col).upper().replace('_', ' ')
        hdr_cells[i].paragraphs[0].runs[0].bold = True
        
    for item in lista_dados:
        row_cells = table.add_row().cells 
        for i, col in enumerate(columns):
            try:
                valor = getattr(item, col, '-') 
                row_cells[i].text = str(valor)
            except:
                row_cells[i].text = '-'
                
    document.add_paragraph()

def renderizar_texto_inline(node, paragraph, relatorio):
    node_type = node.get('type')
    
    if node_type == 'text':
        texto = node.get('text', '')
        run = paragraph.add_run(texto)
        for mark in node.get('marks', []):
            m_type = mark.get('type')
            if m_type == 'strong': run.bold = True
            elif m_type == 'em': run.italic = True
            elif m_type == 'underline': run.underline = True
            
    elif node_type == 'variable':
        ref = node.get('attrs', {}).get('ref', '')
        valor_real = obter_valor_real(relatorio, ref)
        
        run = paragraph.add_run(valor_real)
        run.font.color.rgb = RGBColor(107, 33, 168)
        run.bold = True

def renderizar_bloco(node, document, relatorio):
    node_type = node.get('type')
    
    if node_type == 'paragraph':
        p = document.add_paragraph()
        for child in node.get('content', []):
            renderizar_texto_inline(child, p, relatorio)
            
    elif node_type == 'heading':
        level = node.get('attrs', {}).get('level', 1)
        h = document.add_heading('', level=level)
        for child in node.get('content', []):
            renderizar_texto_inline(child, h, relatorio)
            
    elif node_type == 'table_ref':
        renderizar_tabela(node, document, relatorio)
        
    elif node_type == 'image_ref':
        renderizar_imagem(node, document, relatorio)
        
    elif node_type == 'custom_text':
        for child in node.get('content', []):
            renderizar_bloco(child, document, relatorio)
            
    elif node_type == 'blockquote':
        p = document.add_paragraph(style='Intense Quote')
        for child in node.get('content', []):
            renderizar_texto_inline(child, p, relatorio)

def gerar_relatorio_docx(relatorio, output_path):
    document = Document()
    doc_json = relatorio.metadados
    
    if doc_json and isinstance(doc_json, dict) and doc_json.get('type') == 'doc':
        for node in doc_json.get('content', []):
            renderizar_bloco(node, document, relatorio)
    else:
        document.add_paragraph("Documento sem conteúdo JSON salvo.")
            
    document.save(output_path)


def gerar_url_pre_assinada_s3(caminho_do_arquivo, expiracao_em_segundos=300):
    s3_client = boto3.client(
        's3',
        aws_access_key_id=settings.AWS_ACCESS_KEY_ID,
        aws_secret_access_key=settings.AWS_SECRET_ACCESS_KEY,
        region_name=settings.AWS_S3_REGION_NAME
    )

    try:
        response = s3_client.generate_presigned_url(
            'get_object',
            Params={
                'Bucket': settings.AWS_STORAGE_BUCKET_NAME,
                'Key': caminho_do_arquivo
            },
            ExpiresIn=expiracao_em_segundos
        )
    except ClientError as e:
        print(f"Erro ao gerar link S3: {e}")
        return None

    return response